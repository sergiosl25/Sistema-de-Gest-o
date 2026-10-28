# -*- coding: utf-8 -*-
"""Relatórios de Alves Personalizados — Flet 0.25.x, somente leitura."""
import calendar
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread
from urllib.parse import quote
import hashlib
import flet as ft
from firebase_relatorios import FirebaseRelatorios

AZUL = '#2563EB'


def dinheiro(centavos):
    return 'R$ {:,.2f}'.format(centavos / 100).replace(',', 'X').replace('.', ',').replace('X', '.')


def centavos_de(valor):
    try:
        return int((Decimal(str(valor)) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    except Exception:
        return 0


def data_utc(valor):
    try:
        data = datetime.fromisoformat(str(valor).replace('Z', '+00:00'))
        if data.tzinfo is None:
            data = data.replace(tzinfo=timezone.utc)
        return data.astimezone(timezone.utc)
    except (ValueError, TypeError):
        return None


class TelaRelatorios:
    OPCOES = ['Vendas', 'Produtos', 'Clientes', 'Produção', 'Entregas', 'Financeiro']

    def __init__(self, page, sessao):
        self.page = page
        self.api = FirebaseRelatorios(sessao)
        self.dados = {'pedidos': [], 'movimentos': [], 'clientes': [], 'produtos': []}
        self.linhas = []
        self.colunas = []
        self.titulo = ''
        hoje = datetime.now(timezone.utc)
        self.tipo = ft.Dropdown(label='Relatório', width=195, value='Vendas',
                                options=[ft.dropdown.Option(s) for s in self.OPCOES], on_change=self.atualizar)
        self.periodo = ft.Dropdown(label='Período', width=190, value='mes', options=[
            ft.dropdown.Option('mes', 'Mês'), ft.dropdown.Option('ano', 'Ano'),
            ft.dropdown.Option('todos', 'Todo o histórico'), ft.dropdown.Option('personalizado', 'Personalizado')],
            on_change=self.atualizar)
        self.mes = ft.Dropdown(label='Mês', width=130, value=str(hoje.month),
                               options=[ft.dropdown.Option(str(i), str(i).zfill(2)) for i in range(1, 13)],
                               on_change=self.atualizar)
        self.ano = ft.TextField(label='Ano', width=95, value=str(hoje.year), on_change=self.atualizar)
        self.inicio = ft.TextField(label='Início AAAA-MM-DD', width=180, on_change=self.atualizar)
        self.fim = ft.TextField(label='Fim AAAA-MM-DD', width=180, on_change=self.atualizar)
        self.mensagem = ft.Text('')
        self.indicadores = ft.Row(wrap=True, spacing=12)
        self.grafico = ft.Column(spacing=6)
        self.tabela_area = ft.Column()
        self.conteudo = ft.Column([
            ft.Text('Relatórios', size=25, weight=ft.FontWeight.BOLD),
            ft.Text('Consultas e exportações • datas em UTC • sem alteração no Firebase'),
            ft.Row([self.tipo, self.periodo, self.mes, self.ano, self.inicio, self.fim], wrap=True),
            ft.Row([
                ft.ElevatedButton('Atualizar dados', icon=ft.icons.REFRESH, on_click=self.carregar),
                ft.ElevatedButton('Exportar Excel', icon=ft.icons.TABLE_CHART,
                                  bgcolor='#15803D', color='white', on_click=lambda e: self.exportar('xlsx')),
                ft.ElevatedButton('Exportar PDF', icon=ft.icons.PICTURE_AS_PDF,
                                  bgcolor=AZUL, color='white', on_click=lambda e: self.exportar('pdf')),
            ], wrap=True),
            self.mensagem, self.indicadores, ft.Divider(),
            ft.Text('Resumo visual', size=19, weight=ft.FontWeight.BOLD), self.grafico,
            ft.Divider(), ft.Text('Detalhamento', size=19, weight=ft.FontWeight.BOLD),
            self.tabela_area,
            ft.Text('Vendas e pedidos: somente registros personalizados. Financeiro: somente caixaMovimentos. '
                    'O cartão de crédito é considerado integralmente na data do lançamento.', size=12)
        ], spacing=12, scroll=ft.ScrollMode.AUTO, expand=True)

    def construir(self):
        return self.conteudo

    def intervalo(self):
        agora = datetime.now(timezone.utc)
        modo = self.periodo.value
        if modo == 'todos':
            return None, None
        if modo == 'personalizado':
            try:
                a = datetime.strptime((self.inicio.value or '').strip(), '%Y-%m-%d').replace(tzinfo=timezone.utc)
                b = datetime.strptime((self.fim.value or '').strip(), '%Y-%m-%d').replace(tzinfo=timezone.utc)
            except ValueError:
                raise ValueError('Informe as duas datas no formato AAAA-MM-DD.')
            if b < a:
                raise ValueError('A data final deve ser igual ou posterior à inicial.')
            return a, b + timedelta(days=1)
        try:
            ano = int(self.ano.value)
            mes = int(self.mes.value)
            if not (2000 <= ano <= 9998 and 1 <= mes <= 12):
                raise ValueError()
        except (TypeError, ValueError):
            raise ValueError('Informe ano e mês válidos.')
        if modo == 'ano':
            return datetime(ano, 1, 1, tzinfo=timezone.utc), datetime(ano + 1, 1, 1, tzinfo=timezone.utc)
        a = datetime(ano, mes, 1, tzinfo=timezone.utc)
        b = datetime(ano + 1, 1, 1, tzinfo=timezone.utc) if mes == 12 else datetime(ano, mes + 1, 1, tzinfo=timezone.utc)
        return a, b

    def filtrar(self, registros):
        a, b = self.intervalo()
        saida = []
        for r in registros:
            data = data_utc(r.get('criadoEm'))
            if a is None or (data is not None and a <= data < b):
                saida.append(r)
        return saida

    def carregar(self, e=None):
        try:
            self.mensagem.value = 'Consultando Firebase...'
            self.page.update()
            self.dados = self.api.carregar_dados()
            self.atualizar()
        except Exception as erro:
            self.mensagem.value = 'Erro ao carregar relatórios: ' + str(erro)
            self.page.update()

    def card(self, nome, valor):
        return ft.Container(bgcolor='white', border_radius=12, padding=16, width=215,
                            content=ft.Column([ft.Text(nome, color='#64748B'),
                                               ft.Text(str(valor), size=19, weight=ft.FontWeight.BOLD)], tight=True))

    def atualizar(self, e=None):
        try:
            pedidos = self.filtrar(self.dados['pedidos'])
            movs = self.filtrar(self.dados['movimentos'])
            tipo = self.tipo.value
            linhas = []
            cards = []
            colunas = []
            if tipo == 'Vendas':
                colunas = ['Data UTC', 'Cliente', 'Pedido', 'Itens', 'Valor', 'Entrega']
                for p in pedidos:
                    linhas.append([str(p.get('criadoEm', ''))[:10], p.get('clienteNome', ''), p.get('id', ''),
                                   len(p.get('itens') or []), dinheiro(centavos_de(p.get('total', 0))),
                                   p.get('statusEntrega', 'Pendente')])
                total = sum(centavos_de(p.get('total', 0)) for p in pedidos)
                cards = [('Pedidos', len(pedidos)), ('Total dos pedidos', dinheiro(total)),
                         ('Ticket médio', dinheiro(round(total / len(pedidos))) if pedidos else dinheiro(0))]
            elif tipo == 'Produtos':
                colunas = ['Produto', 'Quantidade vendida', 'Valor dos itens']
                mapa = defaultdict(lambda: [0, 0])
                for p in pedidos:
                    for item in p.get('itens') or []:
                        nome = item.get('produtoNome') or 'Sem nome'
                        mapa[nome][0] += int(item.get('quantidade') or 0)
                        mapa[nome][1] += centavos_de(item.get('subtotal') or 0)
                linhas = [[k, v[0], dinheiro(v[1])] for k, v in sorted(mapa.items(), key=lambda kv: kv[1][0], reverse=True)]
                cards = [('Produtos diferentes', len(mapa)), ('Unidades vendidas', sum(v[0] for v in mapa.values()))]
            elif tipo == 'Clientes':
                colunas = ['Cliente', 'Pedidos', 'Total comprado']
                mapa = defaultdict(lambda: [0, 0])
                for p in pedidos:
                    nome = p.get('clienteNome') or 'Sem nome'
                    mapa[nome][0] += 1
                    mapa[nome][1] += centavos_de(p.get('total') or 0)
                linhas = [[k, v[0], dinheiro(v[1])] for k, v in sorted(mapa.items(), key=lambda kv: kv[1][1], reverse=True)]
                cards = [('Clientes compradores', len(mapa)), ('Pedidos', len(pedidos))]
            elif tipo == 'Produção':
                colunas = ['Status de produção', 'Itens', 'Unidades']
                mapa = defaultdict(lambda: [0, 0])
                for p in pedidos:
                    for item in p.get('itens') or []:
                        status = item.get('statusProducao') or 'Não informado'
                        mapa[status][0] += 1
                        mapa[status][1] += int(item.get('quantidade') or 0)
                linhas = [[k, v[0], v[1]] for k, v in sorted(mapa.items())]
                cards = [('Itens de pedidos', sum(v[0] for v in mapa.values())),
                         ('Unidades', sum(v[1] for v in mapa.values()))]
            elif tipo == 'Entregas':
                colunas = ['Status de entrega', 'Pedidos']
                mapa = defaultdict(int)
                for p in pedidos:
                    entrega = p.get('entrega') or {}
                    status = entrega.get('status') or p.get('statusEntrega') or 'Pendente'
                    mapa[status] += 1
                linhas = [[k, v] for k, v in sorted(mapa.items())]
                cards = [('Pedidos', len(pedidos)), ('Entregues', sum(v for k, v in mapa.items() if k in ('Entregue', 'Concluído')))]
            else:
                colunas = ['Data UTC', 'Tipo', 'Origem', 'Descrição', 'Forma', 'Valor']
                for m in movs:
                    linhas.append([str(m.get('criadoEm', ''))[:10], m.get('tipo', ''), m.get('origem', ''),
                                   m.get('descricao', ''), m.get('forma', ''), dinheiro(int(m.get('centavos') or 0))])
                entradas = sum(int(m.get('centavos') or 0) for m in movs if m.get('tipo') == 'Entrada')
                saidas = sum(int(m.get('centavos') or 0) for m in movs if m.get('tipo') == 'Saída')
                vendas = sum(int(m.get('centavos') or 0) for m in movs if m.get('tipo') == 'Entrada' and m.get('origem') == 'Venda')
                cards = [('Faturamento recebido', dinheiro(vendas)), ('Entradas', dinheiro(entradas)),
                         ('Despesas', dinheiro(saidas)), ('Resultado', dinheiro(entradas - saidas))]
            self.colunas, self.linhas, self.titulo = colunas, linhas, tipo
            self.indicadores.controls = [self.card(k, v) for k, v in cards]
            self.grafico.controls.clear()
            if linhas:
                valores = []
                if tipo == 'Vendas':
                    por_mes = defaultdict(int)
                    for p in pedidos:
                        por_mes[str(p.get('criadoEm', ''))[:7]] += centavos_de(p.get('total', 0))
                    valores = [(k, dinheiro(v), v) for k, v in sorted(por_mes.items())]
                elif tipo == 'Financeiro':
                    por_mes = defaultdict(int)
                    for m in movs:
                        chave = str(m.get('criadoEm', ''))[:7]
                        por_mes[chave] += int(m.get('centavos') or 0) * (1 if m.get('tipo') == 'Entrada' else -1)
                    valores = [(k, dinheiro(v), abs(v)) for k, v in sorted(por_mes.items())]
                elif tipo in ('Produtos', 'Clientes'):
                    valores = [(str(l[0]), str(l[1]), int(l[1])) for l in linhas[:10]]
                else:
                    valores = [(str(l[0]), str(l[1]), int(l[1])) for l in linhas]
                maximo = max((v[2] for v in valores), default=1) or 1
                for rotulo, texto, valor in valores[:18]:
                    self.grafico.controls.append(ft.Row([
                        ft.Text(rotulo[:30], width=190),
                        ft.Container(bgcolor=AZUL, height=16, width=max(3, int(240 * valor / maximo)), border_radius=3),
                        ft.Text(texto)
                    ], wrap=True))
            else:
                self.grafico.controls.append(ft.Text('Sem dados para o período.'))
            self.tabela_area.controls = [ft.Row([ft.DataTable(
                columns=[ft.DataColumn(ft.Text(c)) for c in colunas],
                rows=[ft.DataRow(cells=[ft.DataCell(ft.Text(str(x))) for x in linha]) for linha in linhas[:250]]
            )], scroll=ft.ScrollMode.AUTO)] if linhas else [ft.Text('Nenhum registro encontrado.')]
            self.mensagem.value = '{} registro(s) no relatório. Exibição limitada a 250 linhas; exportação inclui todas.'.format(len(linhas))
        except Exception as erro:
            self.mensagem.value = 'Erro no filtro: ' + str(erro)
        self.page.update()

    def gerar_excel(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = self.titulo[:31]
        ws.append(['ALVES PERSONALIZADOS — ' + self.titulo])
        ws.append(['Gerado em UTC: ' + datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')])
        ws.append([])
        ws.append(self.colunas)
        for linha in self.linhas:
            ws.append([str(v) if not isinstance(v, (int, float)) else v for v in linha])
        for cell in ws[4]:
            cell.fill = PatternFill('solid', fgColor='2563EB')
            cell.font = Font(bold=True, color='FFFFFF')
        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = 'A4:{}{}'.format(get_column_letter(len(self.colunas)), max(4, ws.max_row))
        for i, nome in enumerate(self.colunas, 1):
            tamanho = max([len(str(nome))] + [len(str(l[i - 1])) for l in self.linhas[:500]]) if self.linhas else len(nome)
            ws.column_dimensions[get_column_letter(i)].width = min(55, max(14, tamanho + 3))
        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def gerar_pdf(self):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
        from reportlab.lib.enums import TA_LEFT
        from xml.sax.saxutils import escape
        # Compatibilidade com builds de Python/OpenSSL que não aceitam
        # hashlib.md5(usedforsecurity=False), utilizado pelo ReportLab.
        import reportlab.pdfbase.pdfdoc as pdfdoc
        md5_original = pdfdoc.md5
        def md5_compativel(*args, **kwargs):
            kwargs.pop('usedforsecurity', None)
            return hashlib.md5(*args, **kwargs)
        pdfdoc.md5 = md5_compativel
        output = BytesIO()
        pagina = landscape(A4)
        doc = SimpleDocTemplate(output, pagesize=pagina, leftMargin=28, rightMargin=28,
                                topMargin=30, bottomMargin=30)
        estilos = getSampleStyleSheet()
        elementos = [Paragraph('ALVES PERSONALIZADOS — ' + escape(self.titulo), estilos['Heading1']),
                     Paragraph('Gerado em UTC: ' + datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M'), estilos['Normal']),
                     Spacer(1, 14)]
        # Quebra texto dentro das células para evitar extrapolar a página.
        estilo = estilos['BodyText']
        estilo.fontSize = 7
        estilo.leading = 9
        linhas = [[Paragraph(escape(str(x)), estilo) for x in self.colunas]]
        linhas += [[Paragraph(escape(str(x)), estilo) for x in linha] for linha in self.linhas]
        if not self.linhas:
            linhas.append([Paragraph('Nenhum registro', estilo)] + [''] * (len(self.colunas) - 1))
        largura = pagina[0] - 56
        n = len(self.colunas)
        tabela = LongTable(linhas, repeatRows=1, colWidths=[largura / n] * n, hAlign='LEFT')
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('LINEBELOW', (0, 0), (-1, 0), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        elementos.append(tabela)
        try:
            doc.build(elementos)
        finally:
            pdfdoc.md5 = md5_original
        return output.getvalue()

    def exportar(self, formato):
        try:
            self.atualizar()
            conteudo = self.gerar_excel() if formato == 'xlsx' else self.gerar_pdf()
            nome = 'alves_{}_{}.{}'.format(
                self.titulo.lower().replace('ç', 'c').replace('ã', 'a'),
                datetime.now().strftime('%Y%m%d_%H%M%S'), formato)

            # Flet 0.25.x Web: FilePicker.save_file não recebe src_bytes.
            # Um servidor HTTP local entrega os bytes reais ao navegador.
            mime = ('application/vnd.openxmlformats-officedocument.'
                    'spreadsheetml.sheet' if formato == 'xlsx' else 'application/pdf')
            class DownloadHandler(BaseHTTPRequestHandler):
                def do_GET(self):
                    self.send_response(200)
                    self.send_header('Content-Type', mime)
                    self.send_header('Content-Disposition',
                                     'attachment; filename="{}"'.format(nome))
                    self.send_header('Content-Length', str(len(conteudo)))
                    self.send_header('Cache-Control', 'no-store')
                    self.end_headers()
                    self.wfile.write(conteudo)
                def log_message(self, *args):
                    pass

            servidor = ThreadingHTTPServer(('127.0.0.1', 0), DownloadHandler)
            porta = servidor.server_address[1]
            Thread(target=servidor.serve_forever, daemon=True).start()
            self.page.launch_url('http://127.0.0.1:{}/download'.format(porta))
            self.mensagem.value = 'Download solicitado: ' + nome
        except ImportError as erro:
            self.mensagem.value = ('Biblioteca ausente: ' + str(erro) +
                                   '. Instale openpyxl e reportlab.')
        except Exception as erro:
            self.mensagem.value = 'Erro ao exportar: ' + str(erro)
        self.page.update()
